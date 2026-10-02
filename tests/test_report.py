import csv
import json
from types import SimpleNamespace
from contextlib import closing

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from kol.cli import app
from kol.fraud import detect_fraud
from kol.metrics import cpm, parse_price
from kol.report import ChannelRow, export_xlsx, select_rows
from kol.sources.telegram import _fetch_channel
from kol.sources.youtube import _fetch_youtube_channel


def test_report_filters_apply_to_exports_and_unicode(tmp_path):
    source = tmp_path / "channels.json"
    source.write_text(json.dumps([
        {"handle": "@дорого", "cpm": 20, "er_pct": 2},
        {"handle": "@канал", "cpm": 3, "er_pct": 5, "frequency_per_week": 2},
        {"handle": "@unknown", "cpm": None, "er_pct": 6},
    ], ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "report.csv"
    result = CliRunner().invoke(app, ["report", str(source), "--max-cpm", "5", "--csv", str(output)])
    assert result.exit_code == 0, result.output
    with output.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["handle"] for row in rows] == ["@канал"]
    assert rows[0]["frequency"] == "2.0"


def test_sort_missing_last_and_preserve_input():
    rows = [ChannelRow("missing"), ChannelRow("high", cpm=10), ChannelRow("low", cpm=2)]
    assert [r.handle for r in select_rows(rows, "cpm")] == ["low", "high", "missing"]
    assert rows[0].handle == "missing"


def test_xlsx_empty_headers_and_untrusted_text(tmp_path):
    output = tmp_path / "report.xlsx"
    export_xlsx([], str(output))
    with closing(load_workbook(output)) as workbook:
        assert workbook.active.cell(1, 1).value == "Handle"
    export_xlsx([ChannelRow("@foo", title="=1+1")], str(output))
    with closing(load_workbook(output)) as workbook:
        assert workbook.active.cell(2, 3).value == "=1+1"
        assert workbook.active.cell(2, 3).data_type == "s"


@pytest.mark.parametrize("value", [float("nan"), float("inf"), "inf", "NaN", True])
def test_invalid_prices_are_not_metrics(value):
    assert parse_price(value) is None
    if isinstance(value, float):
        assert cpm(value, 1000) is None


def test_zero_reach_is_fraud_signal_and_nan_is_ignored():
    flags = detect_fraud(subscribers=10000, reach=0)
    assert any(flag.code == "LOW_REACH_RATIO" and flag.severity == "high" for flag in flags)
    assert detect_fraud(subscribers=10000, reach=float("nan"), er_pct=float("inf")) == []


@pytest.mark.asyncio
async def test_telegram_service_messages_do_not_crash_collection():
    responses = iter([
        SimpleNamespace(peer="peer", chats=[SimpleNamespace(title="Test", participants_count=None)]),
        SimpleNamespace(full_chat=SimpleNamespace(participants_count=1200)),
        SimpleNamespace(messages=[SimpleNamespace(), SimpleNamespace(views=20)]),
    ])
    async def client(request):
        return next(responses)
    result = await _fetch_channel(client, "@foo", 50, lambda *a: a, lambda **kw: kw)
    assert result.total_views_last_n == 20
    assert result.subscribers == 1200


@pytest.mark.asyncio
async def test_youtube_unknown_analytics_are_not_zero():
    payload = {"header": {"c4TabbedHeaderRenderer": {"title": "Channel", "subscriberCountText": {"simpleText": "1.2K subscribers"}}}}
    class Client:
        async def get(self, url):
            return SimpleNamespace(text="<script>var ytInitialData =\n" + json.dumps(payload) + ";\n</script>", raise_for_status=lambda: None)
    result = await _fetch_youtube_channel(Client(), "@channel", 10)
    assert result.title == "Channel"
    assert result.subscribers == 1200
    assert result.avg_views is None and result.er_pct is None
    assert result.platform == "yt"
